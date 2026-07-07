"""Safely add/update one row in the user's existing Excel tracking sheet.

Loads the workbook in place (openpyxl, not pandas) so existing rows, styles, and any
other sheets in the file are left untouched. Re-running for the same version_no updates
that row instead of appending a duplicate. Backs up before writing and restores on any
failure, since this file is the user's own record and must never end up corrupted or
partially written.
"""

from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

import openpyxl

from doc_update_tool.config import TrackingSheetConfig

_HYPERLINK_FIELDS = {"email_path", "bilingual_file_path", "diff_report_path"}


class TrackingSheetError(RuntimeError):
    pass


def upsert_row(config: TrackingSheetConfig, version_no: int, row_data: dict) -> None:
    path = config.path
    if not path.exists():
        raise TrackingSheetError(
            f"追蹤表不存在: {path}。請先建立好既有追蹤表（含表頭），本工具只會新增/更新資料列。"
        )

    backup_path = path.with_suffix(path.suffix + ".bak")
    shutil.copy2(path, backup_path)

    try:
        wb = openpyxl.load_workbook(path)
        ws = wb[config.sheet_name]

        version_col = config.columns["version_no"]
        target_row = _find_row_by_version(ws, version_col, config.header_row, version_no)

        if target_row is None:
            target_row = ws.max_row + 1
            if ws.max_row > config.header_row:
                _copy_row_style(ws, ws.max_row, target_row)

        for field, value in row_data.items():
            col_letter = config.columns.get(field)
            if col_letter is None:
                continue  # unrecognized field name, ignore rather than fail the whole upsert
            cell = ws[f"{col_letter}{target_row}"]
            if field in _HYPERLINK_FIELDS and value:
                value_path = Path(value)
                cell.value = value_path.name
                cell.hyperlink = str(value_path)
            else:
                cell.value = value

        wb.save(path)
        openpyxl.load_workbook(path)  # sanity check: the file we just wrote must still open
    except Exception:
        shutil.copy2(backup_path, path)
        raise
    finally:
        backup_path.unlink(missing_ok=True)


def read_row(config: TrackingSheetConfig, version_no: int) -> dict[str, object] | None:
    if not config.path.exists():
        return None
    wb = openpyxl.load_workbook(config.path, data_only=True)
    ws = wb[config.sheet_name]
    row = _find_row_by_version(ws, config.columns["version_no"], config.header_row, version_no)
    if row is None:
        return None
    return {field: ws[f"{col}{row}"].value for field, col in config.columns.items()}


def _find_row_by_version(ws, version_col: str, header_row: int, version_no: int) -> int | None:
    for row in range(header_row + 1, ws.max_row + 1):
        cell_value = ws[f"{version_col}{row}"].value
        if cell_value is not None and str(cell_value).strip() == str(version_no):
            return row
    return None


def _copy_row_style(ws, src_row: int, dst_row: int) -> None:
    for col_idx in range(1, ws.max_column + 1):
        src_cell = ws.cell(row=src_row, column=col_idx)
        dst_cell = ws.cell(row=dst_row, column=col_idx)
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.alignment = copy(src_cell.alignment)
