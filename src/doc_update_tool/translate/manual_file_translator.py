"""Placeholder translation step: produces a worklist xlsx for a human to fill in, and
reads the filled-in file back. This is deliberately not the `Translator` Protocol —
translation here spans two CLI invocations (`process` then `apply-translations`) with a
human editing a file in between, not a single synchronous call. Swapping in an
API/LLM-backed Translator later replaces this module without changing the pipeline.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from doc_update_tool.models import TranslationItem, TranslationResult

_HEADERS = ["change_id", "中文原文", "context", "英文翻譯"]


class ManualFileTranslator:
    def write_worklist(self, items: list[TranslationItem], path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "待翻譯清單"
        ws.append(_HEADERS)
        for item in items:
            ws.append([item.change_id, item.zh_text, item.context or "", ""])
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    def read_results(self, path: Path) -> list[TranslationResult]:
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        results: list[TranslationResult] = []
        missing: list[str] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or row[0] is None:
                continue
            change_id, _zh_text, _context, en_text = (list(row) + [None] * 4)[:4]
            if not en_text or not str(en_text).strip():
                missing.append(str(change_id))
                continue
            results.append(
                TranslationResult(change_id=str(change_id), en_text=str(en_text).strip(), source="manual")
            )
        if missing:
            raise ValueError(
                f"待翻譯清單仍有 {len(missing)} 筆未填寫英文翻譯，請填完再執行 apply-translations: "
                + ", ".join(missing)
            )
        return results
