from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from doc_update_tool.config import TrackingSheetConfig
from doc_update_tool.tracking.tracking_sheet import TrackingSheetError, read_row, upsert_row

_COLUMNS = {
    "version_no": "A",
    "received_date": "B",
    "sender": "C",
    "email_path": "D",
    "bilingual_file_path": "E",
    "diff_report_path": "F",
    "update_summary": "G",
    "is_first_version": "H",
    "overseas_confirm_status": "I",
    "overseas_confirm_date": "J",
    "overseas_confirm_method": "K",
    "customer_confirm_status": "L",
    "customer_confirm_date": "M",
    "customer_confirm_method": "N",
    "status": "O",
    "notes": "P",
}


def _make_sheet(tmp_path) -> TrackingSheetConfig:
    path = tmp_path / "追蹤表.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "追蹤"
    for field, col in _COLUMNS.items():
        ws[f"{col}1"] = field
        ws[f"{col}1"].font = Font(bold=True)
    # one pre-existing row so upsert has a style to copy for the next appended row
    ws["A2"] = 1
    ws["G2"] = "首次提供文件"
    for col in _COLUMNS.values():
        ws[f"{col}2"].font = Font(italic=True)
    wb.save(path)
    return TrackingSheetConfig(path=path, sheet_name="追蹤", header_row=1, columns=dict(_COLUMNS))


def test_upsert_appends_new_row_and_copies_style(tmp_path):
    config = _make_sheet(tmp_path)
    upsert_row(
        config,
        version_no=2,
        row_data={
            "version_no": 2,
            "update_summary": "第二次更新",
            "bilingual_file_path": str(tmp_path / "archive" / "v02" / "bilingual.xlsx"),
        },
    )

    wb = openpyxl.load_workbook(config.path)
    ws = wb["追蹤"]
    assert ws.max_row == 3
    assert ws["A3"].value == 2
    assert ws["G3"].value == "第二次更新"
    assert ws["A3"].font.italic is True  # style copied from row 2

    # path field becomes a hyperlink with just the filename shown, not the full path
    assert ws["E3"].value == "bilingual.xlsx"
    assert ws["E3"].hyperlink.target == str(tmp_path / "archive" / "v02" / "bilingual.xlsx")


def test_upsert_same_version_updates_row_instead_of_duplicating(tmp_path):
    config = _make_sheet(tmp_path)
    upsert_row(config, version_no=2, row_data={"version_no": 2, "status": "待確認"})
    upsert_row(config, version_no=2, row_data={"status": "已完成"})

    wb = openpyxl.load_workbook(config.path)
    ws = wb["追蹤"]
    assert ws.max_row == 3  # still just header + v1 + v2, no duplicate v2 row
    assert ws["O3"].value == "已完成"


def test_read_row_returns_none_when_version_missing(tmp_path):
    config = _make_sheet(tmp_path)
    assert read_row(config, version_no=99) is None


def test_read_row_returns_existing_values(tmp_path):
    config = _make_sheet(tmp_path)
    row = read_row(config, version_no=1)
    assert row["update_summary"] == "首次提供文件"


def test_upsert_raises_if_file_missing(tmp_path):
    config = TrackingSheetConfig(
        path=tmp_path / "does_not_exist.xlsx", sheet_name="追蹤", header_row=1, columns=dict(_COLUMNS)
    )
    with pytest.raises(TrackingSheetError):
        upsert_row(config, version_no=1, row_data={"version_no": 1})


def test_backup_file_is_cleaned_up_after_successful_upsert(tmp_path):
    config = _make_sheet(tmp_path)
    upsert_row(config, version_no=2, row_data={"version_no": 2})
    assert not config.path.with_suffix(".xlsx.bak").exists()
