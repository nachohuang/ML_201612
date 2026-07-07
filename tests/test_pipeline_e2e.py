"""End-to-end golden path: a zh-only v1 document, then a v2 document with one modified
row and one added row, driven through process -> (fill worklist) -> apply-translations."""

from pathlib import Path

import openpyxl

from doc_update_tool.archive.manager import ArchiveManager
from doc_update_tool.config import ClientConfig, ExcelMapping, Folders, TrackingSheetConfig, WordMapping
from doc_update_tool.pipeline import first_version, update_version
from doc_update_tool.pipeline.apply import apply_translations
from doc_update_tool.translate.manual_file_translator import ManualFileTranslator

_TRACKING_COLUMNS = {
    "version_no": "A",
    "received_date": "B",
    "sender": "C",
    "email_path": "D",
    "bilingual_file_path": "E",
    "diff_report_path": "F",
    "update_summary": "G",
    "is_first_version": "H",
    "overseas_confirm_status": "I",
    "overseas_confirm_date": "J",
    "overseas_confirm_method": "K",
    "customer_confirm_status": "L",
    "customer_confirm_date": "M",
    "customer_confirm_method": "N",
    "status": "O",
    "notes": "P",
}


def _make_config(tmp_path: Path) -> ClientConfig:
    tracking_path = tmp_path / "追蹤表.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "追蹤"
    for field, col in _TRACKING_COLUMNS.items():
        ws[f"{col}1"] = field
    wb.save(tracking_path)

    folders = Folders(incoming=tmp_path / "incoming", archive=tmp_path / "archive", output=tmp_path / "output")
    folders.incoming.mkdir()

    return ClientConfig(
        client_id="test_client",
        folders=folders,
        tracking_sheet=TrackingSheetConfig(
            path=tracking_path, sheet_name="追蹤", header_row=1, columns=dict(_TRACKING_COLUMNS)
        ),
        document_format="excel",
        excel_mapping=ExcelMapping(
            sheet_name="auto", header_row=1, key_columns=["A"], zh_columns=["C"], en_columns=["D"]
        ),
        word_mapping=WordMapping(),
    )


def _write_v1_customer_doc(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "編號"
    ws["C1"] = "需求說明"
    ws["D1"] = "English"
    ws["A2"], ws["C2"] = 1, "第一項需求"
    ws["A3"], ws["C3"] = 2, "第二項需求"
    wb.save(path)


def _fill_worklist(worklist_path: Path, translations: dict[str, str]) -> None:
    wb = openpyxl.load_workbook(worklist_path)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        change_id = row[0].value
        if change_id in translations:
            row[3].value = translations[change_id]
    wb.save(worklist_path)


def test_v1_then_v2_golden_path(tmp_path):
    config = _make_config(tmp_path)
    archive = ArchiveManager(config.folders.incoming, config.folders.archive, config.folders.output)

    # --- v1: zh-only document, no previous version ---
    assert archive.is_first_version()
    v1_doc = config.folders.incoming / "customer_v1.xlsx"
    _write_v1_customer_doc(v1_doc)

    worklist_path = first_version.process_first_version(config, archive)
    _fill_worklist(
        worklist_path,
        {"Sheet!D2": "Item one", "Sheet!D3": "Item two"},
    )
    v1_dir = apply_translations(config, archive, worklist_path)

    assert archive.existing_versions() == [1]
    v1_bilingual = v1_dir / "bilingual.xlsx"
    wb = openpyxl.load_workbook(v1_bilingual)
    ws = wb.active
    assert ws["D2"].value == "Item one"
    assert ws["D3"].value == "Item two"

    tracking_wb = openpyxl.load_workbook(config.tracking_sheet.path)
    tracking_ws = tracking_wb["追蹤"]
    assert tracking_ws["A2"].value == 1
    assert tracking_ws["H2"].value == "是"
    assert tracking_ws["O2"].value == "待確認"
    assert tracking_ws["E2"].value == "bilingual.xlsx"
    assert tracking_ws["E2"].hyperlink.target == str(v1_bilingual)

    v1_doc.unlink()  # simulate incoming/ being cleared out after processing

    # --- v2: customer edited the bilingual file — old English + new zh mixed together ---
    v2_doc = config.folders.incoming / "customer_v2.xlsx"
    wb2 = openpyxl.load_workbook(v1_bilingual)
    ws2 = wb2.active
    ws2["C3"] = "第二項需求（已修改）"  # modified
    ws2["A4"], ws2["C4"] = 3, "第三項需求（新增）"  # added
    wb2.save(v2_doc)

    assert not archive.is_first_version()
    worklist_path_v2 = update_version.process_update_version(config, archive)

    worklist_wb = openpyxl.load_workbook(worklist_path_v2)
    worklist_rows = list(worklist_wb.active.iter_rows(min_row=2, values_only=True))
    change_ids = {row[0] for row in worklist_rows}
    # only the modified + added rows need translation; unmodified row 2 is absent
    assert change_ids == {"Sheet!D3", "Sheet!D4"}

    diff_report_path = config.folders.output.parent / ".working" / "diff_report.xlsx"
    diff_wb = openpyxl.load_workbook(diff_report_path)
    diff_rows = list(diff_wb.active.iter_rows(min_row=2, values_only=True))
    change_types = {row[4] for row in diff_rows}
    assert change_types == {"修改", "新增"}

    _fill_worklist(
        worklist_path_v2,
        {"Sheet!D3": "Item two revised", "Sheet!D4": "Item three"},
    )
    v2_dir = apply_translations(config, archive, worklist_path_v2)

    assert archive.existing_versions() == [1, 2]
    v2_bilingual = v2_dir / "bilingual.xlsx"
    wb3 = openpyxl.load_workbook(v2_bilingual)
    ws3 = wb3.active
    assert ws3["D2"].value == "Item one"  # unchanged row keeps its original translation
    assert ws3["D3"].value == "Item two revised"
    assert ws3["D4"].value == "Item three"

    tracking_wb2 = openpyxl.load_workbook(config.tracking_sheet.path)
    tracking_ws2 = tracking_wb2["追蹤"]
    assert tracking_ws2.max_row == 3  # header + v1 + v2, no duplicate rows
    assert tracking_ws2["A3"].value == 2
