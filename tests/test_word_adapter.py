from pathlib import Path

import docx

from doc_update_tool.config import WordMapping
from doc_update_tool.file_types.word_adapter import WordAdapter, inspect, is_chinese_text
from doc_update_tool.models import SegmentUpdate


def test_is_chinese_text():
    assert is_chinese_text("這是一段中文說明")
    assert not is_chinese_text("This is English text")
    assert not is_chinese_text("")


def test_alternating_paragraphs_extract_and_apply(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("中文段落一")
    doc.add_paragraph("English paragraph one")
    doc.add_paragraph("中文段落二")
    doc.add_paragraph("English paragraph two")
    src = tmp_path / "sample.docx"
    doc.save(src)

    mapping = WordMapping(mode="alternating_paragraphs")
    adapter = WordAdapter(mapping)
    loaded = adapter.load(src)
    segments = adapter.extract_segments(loaded)

    assert len(segments) == 2
    assert segments[0].zh_text == "中文段落一"
    assert segments[0].en_text == "English paragraph one"
    assert segments[0].en_location_id == "para:1"

    adapter.apply_translations(
        loaded, [SegmentUpdate(en_location_id=segments[0].en_location_id, en_text="Revised paragraph one")]
    )
    out = tmp_path / "out.docx"
    adapter.save(loaded, out)

    reloaded = docx.Document(str(out))
    assert reloaded.paragraphs[1].text == "Revised paragraph one"
    assert reloaded.paragraphs[2].text == "中文段落二"  # untouched


def test_alternating_paragraphs_first_version_has_no_english_yet(tmp_path):
    """v1 documents are zh-only — there's nowhere to write English yet, so extraction
    must produce an insert_after target rather than pointing at the zh paragraph itself."""
    doc = docx.Document()
    doc.add_paragraph("中文段落一")
    doc.add_paragraph("中文段落二")
    src = tmp_path / "sample.docx"
    doc.save(src)

    mapping = WordMapping(mode="alternating_paragraphs")
    adapter = WordAdapter(mapping)
    loaded = adapter.load(src)
    segments = adapter.extract_segments(loaded)

    assert len(segments) == 2
    assert segments[0].en_location_id == "para:0:insert_after"
    assert segments[1].en_location_id == "para:1:insert_after"

    # Apply both translations in one call — this exercises the descending-index
    # ordering that keeps insert_after targets valid even after earlier insertions.
    adapter.apply_translations(
        loaded,
        [
            SegmentUpdate(en_location_id=segments[0].en_location_id, en_text="English one"),
            SegmentUpdate(en_location_id=segments[1].en_location_id, en_text="English two"),
        ],
    )
    out = tmp_path / "out.docx"
    adapter.save(loaded, out)

    reloaded = docx.Document(str(out))
    texts = [p.text for p in reloaded.paragraphs]
    assert texts == ["中文段落一", "English one", "中文段落二", "English two"]


def test_same_paragraph_split(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("中文內容|English content")
    src = tmp_path / "sample.docx"
    doc.save(src)

    mapping = WordMapping(mode="same_paragraph_split", split_delimiter="|")
    adapter = WordAdapter(mapping)
    loaded = adapter.load(src)
    segments = adapter.extract_segments(loaded)

    assert len(segments) == 1
    assert segments[0].zh_text == "中文內容"
    assert segments[0].en_text == "English content"

    adapter.apply_translations(
        loaded, [SegmentUpdate(en_location_id=segments[0].en_location_id, en_text="New English content")]
    )
    out = tmp_path / "out.docx"
    adapter.save(loaded, out)

    reloaded = docx.Document(str(out))
    assert reloaded.paragraphs[0].text == "中文內容|New English content"


def test_table_based(tmp_path):
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "中文"
    table.rows[0].cells[1].text = "English"
    table.rows[1].cells[0].text = "第一項"
    table.rows[1].cells[1].text = ""
    table.rows[2].cells[0].text = "第二項"
    table.rows[2].cells[1].text = ""
    src = tmp_path / "sample.docx"
    doc.save(src)

    mapping = WordMapping(mode="table_based", zh_columns=["0"], en_columns=["1"])
    adapter = WordAdapter(mapping)
    loaded = adapter.load(src)
    segments = adapter.extract_segments(loaded)

    assert len(segments) == 2
    assert segments[0].zh_text == "第一項"
    assert segments[0].en_location_id == "table:0:1:1"

    adapter.apply_translations(
        loaded, [SegmentUpdate(en_location_id=segments[0].en_location_id, en_text="Item one")]
    )
    out = tmp_path / "out.docx"
    adapter.save(loaded, out)

    reloaded = docx.Document(str(out))
    assert reloaded.tables[0].rows[1].cells[1].text == "Item one"


def test_inspect_word(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("中文段落")
    doc.add_paragraph("English paragraph")
    src = tmp_path / "sample.docx"
    doc.save(src)

    rows = inspect(src)
    assert rows[0]["is_chinese"] is True
    assert rows[1]["is_chinese"] is False
