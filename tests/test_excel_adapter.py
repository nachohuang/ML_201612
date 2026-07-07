from pathlib import Path

import openpyxl
import pytest

from doc_update_tool.config import ExcelMapping
from doc_update_tool.file_types.excel_adapter import ExcelAdapter, inspect
from doc_update_tool.models import SegmentUpdate


def _build_workbook(path: Path, *, zh_col="C", en_col="D", header_row=1, key_col="A",
                     data_start_row_offset=1, add_formula=False, merge_range=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws[f"A{header_row}"] = "編號"
    ws[f"{zh_col}{header_row}"] = "需求說明"
    ws[f"{en_col}{header_row}"] = "English"
    ws[f"A{header_row}"].font = openpyxl.styles.Font(bold=True)

    row1 = header_row + data_start_row_offset
    ws[f"A{row1}"] = 1
    ws[f"{zh_col}{row1}"] = "第一項需求"
    ws[f"{en_col}{row1}"] = ""

    row2 = row1 + 1
    ws[f"A{row2}"] = 2
    ws[f"{zh_col}{row2}"] = "第二項需求"
    ws[f"{en_col}{row2}"] = ""

    if add_formula:
        ws["Z1"] = "=1+1"

    if merge_range:
        ws.merge_cells(merge_range)

    wb.save(path)
    return path


def test_extract_segments_adjacent_columns(tmp_path):
    path = _build_workbook(tmp_path / "sample.xlsx", zh_col="C", en_col="D")
    mapping = ExcelMapping(sheet_name="auto", header_row=1, key_columns=["A"], zh_columns=["C"], en_columns=["D"])
    adapter = ExcelAdapter(mapping)
    wb = adapter.load(path)
    segments = adapter.extract_segments(wb)

    assert len(segments) == 2
    assert segments[0].zh_text == "第一項需求"
    assert segments[0].row_key == "1"
    assert segments[0].field_name == "需求說明"
    assert segments[0].location_id.endswith("!C2")
    assert segments[0].en_location_id.endswith("!D2")


def test_extract_segments_non_adjacent_columns(tmp_path):
    """Columns need not be next to each other — layout is entirely config-driven."""
    path = _build_workbook(tmp_path / "sample.xlsx", zh_col="B", en_col="F")
    mapping = ExcelMapping(sheet_name="auto", header_row=1, key_columns=["A"], zh_columns=["B"], en_columns=["F"])
    adapter = ExcelAdapter(mapping)
    segments = adapter.extract_segments(adapter.load(path))

    assert len(segments) == 2
    assert segments[0].zh_text == "第一項需求"
    assert segments[0].en_location_id.endswith("!F2")


def test_extract_segments_header_not_on_first_row(tmp_path):
    path = _build_workbook(tmp_path / "sample.xlsx", header_row=3, data_start_row_offset=1)
    mapping = ExcelMapping(sheet_name="auto", header_row=3, key_columns=["A"], zh_columns=["C"], en_columns=["D"])
    adapter = ExcelAdapter(mapping)
    segments = adapter.extract_segments(adapter.load(path))

    assert len(segments) == 2
    assert segments[0].field_name == "需求說明"
    assert segments[0].location_id.endswith("!C4")


def test_apply_translations_round_trip_preserves_rest_of_file(tmp_path):
    src = _build_workbook(tmp_path / "sample.xlsx", add_formula=True, merge_range="F1:G1")
    mapping = ExcelMapping(sheet_name="auto", header_row=1, key_columns=["A"], zh_columns=["C"], en_columns=["D"])
    adapter = ExcelAdapter(mapping)

    wb = adapter.load(src)
    segments = adapter.extract_segments(wb)
    adapter.apply_translations(
        wb,
        [
            SegmentUpdate(en_location_id=segments[0].en_location_id, en_text="Item one"),
            SegmentUpdate(en_location_id=segments[1].en_location_id, en_text="Item two"),
        ],
    )
    out = tmp_path / "out.xlsx"
    adapter.save(wb, out)

    reloaded = openpyxl.load_workbook(out)
    ws = reloaded.active
    assert ws["D2"].value == "Item one"
    assert ws["D3"].value == "Item two"
    assert ws["C2"].value == "第一項需求"  # zh column untouched
    assert ws["Z1"].value == "=1+1"  # unrelated formula untouched
    assert ws["A1"].font.bold is True  # header style preserved
    assert "F1:G1" in [str(r) for r in ws.merged_cells.ranges]  # merged cells preserved


def test_extract_segments_skips_fully_blank_rows(tmp_path):
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "編號"
    ws["C1"] = "需求說明"
    ws["D1"] = "English"
    ws["A2"] = 1
    ws["C2"] = "內容"
    # rows 3-5 intentionally blank (trailing whitespace in the sheet)
    wb.save(path)

    mapping = ExcelMapping(sheet_name="auto", header_row=1, key_columns=["A"], zh_columns=["C"], en_columns=["D"])
    adapter = ExcelAdapter(mapping)
    segments = adapter.extract_segments(adapter.load(path))
    assert len(segments) == 1


def test_inspect_flags_formula_columns(tmp_path):
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "計算欄"
    ws["A2"] = "=1+1"
    wb.save(path)

    cols = inspect(path)
    calc_col = next(c for c in cols if c["header"] == "計算欄")
    assert calc_col["has_formula"] is True
